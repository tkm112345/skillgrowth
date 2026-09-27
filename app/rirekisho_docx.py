from io import BytesIO
from pathlib import Path

from docx.shared import Mm
from docxtpl import DocxTemplate, InlineImage
from sqlmodel import Session

from app.models import RirekishoTemplate
from app.rirekisho_builder import gather_rirekisho_context

# JIS-style rirekisho photo boxes are conventionally about 30mm x 40mm —
# fixed here since there's no per-template UI for it, unlike the resume's
# bullet/table section-format choice.
PHOTO_WIDTH = Mm(30)


def _build_history_subdoc(tpl: DocxTemplate, history: list[dict]):
    subdoc = tpl.new_subdoc()
    table = subdoc.add_table(rows=1, cols=3)
    try:
        table.style = "Table Grid"
    except KeyError:
        pass  # the uploaded template doesn't define this built-in style
    for cell, header in zip(table.rows[0].cells, ["年", "月", "学歴・職歴"]):
        cell.text = header
    for row in history:
        cells = table.add_row().cells
        cells[0].text = row["year"]
        cells[1].text = row["month"]
        cells[2].text = row["label"]
    return subdoc


def _build_certifications_subdoc(tpl: DocxTemplate, certifications: list[dict]):
    subdoc = tpl.new_subdoc()
    for c in certifications:
        suffix = f" ({c['date']})" if c["date"] else ""
        subdoc.add_paragraph(f"• {c['title']}{suffix}")
    return subdoc


def _build_self_pr_subdoc(tpl: DocxTemplate, text: str | None):
    subdoc = tpl.new_subdoc()
    for line in (text or "").split("\n"):
        if line.strip():
            subdoc.add_paragraph(line)
    return subdoc


def render_rirekisho_docx(session: Session, template: RirekishoTemplate) -> bytes:
    """Fill the uploaded .docx template's tags with rirekisho data — no LLM
    call. Unlike the resume's Word export, there's no per-section
    bullet/table format choice; the tag set here is fixed."""
    ctx = gather_rirekisho_context(session)

    tpl = DocxTemplate(template.file_path)
    photo = ""
    if ctx["photo_path"] and Path(ctx["photo_path"]).is_file():
        photo = InlineImage(tpl, ctx["photo_path"], width=PHOTO_WIDTH)

    context = {
        "name": ctx["name"],
        "name_kana": ctx["name_kana"],
        "birthdate": ctx["birthdate"],
        "age": ctx["age"] if ctx["age"] is not None else "",
        "postal_code": ctx["postal_code"],
        "address": ctx["address"],
        "address_kana": ctx["address_kana"],
        "phone": ctx["phone"],
        "email": ctx["email"],
        "photo": photo,
        "history": _build_history_subdoc(tpl, ctx["history"]),
        "certifications": _build_certifications_subdoc(tpl, ctx["certifications"]),
        "self_pr": _build_self_pr_subdoc(tpl, ctx["self_pr"]),
    }
    tpl.render(context)

    buf = BytesIO()
    tpl.save(buf)
    return buf.getvalue()
