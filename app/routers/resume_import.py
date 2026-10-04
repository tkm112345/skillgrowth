from pathlib import Path

import pymupdf
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from pydantic import BaseModel
from sqlmodel import Session, select

from app import llm, services
from app.db import get_session
from app.models import Project, Settings, Skill, SkillLink

router = APIRouter(prefix="/api/resume-import", tags=["resume_import"])

MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024  # 10MB, matches evidence/portfolio upload limits


def _pdf_text(data: bytes) -> str:
    # Text-only: a scanned/image-only PDF yields no text here and falls
    # through to the same "No text content found" error as an empty .docx,
    # rather than silently returning nothing. No OCR fallback — that would
    # mean routing through the vision model instead of the text one, a
    # different extraction path this endpoint doesn't use.
    try:
        with pymupdf.open(stream=data, filetype="pdf") as doc:
            return "\n".join(page.get_text() for page in doc)
    except Exception:
        raise HTTPException(status_code=400, detail="Could not read PDF file")


def _document_text(document) -> str:
    # python-docx's `document.paragraphs` only walks body-level paragraphs —
    # it silently skips any paragraph inside a table. A 職務経歴書's
    # period-by-period work history is commonly laid out as a table (seen
    # firsthand on a real resume while building this feature: two tables,
    # 13 and 4 rows, holding almost the entire project history — reading
    # `.paragraphs` alone would have dropped nearly all of it). Walk the
    # body's child elements directly so tables are read in their original
    # position relative to the surrounding paragraphs, not just paragraphs.
    lines = []
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            text = Paragraph(child, document).text
            if text.strip():
                lines.append(text)
        elif child.tag == qn("w:tbl"):
            for row in Table(child, document).rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    lines.append("\t".join(cells))
    return "\n".join(lines)


@router.post("/extract")
def extract_resume(file: UploadFile = File(...), session: Session = Depends(get_session)) -> dict:
    """Stateless draft extraction: reads the uploaded .docx or .pdf, asks
    the LLM to structure it, and returns the draft as-is. Nothing is
    written to the DB and the file is never saved to disk — the caller
    (the resume-import UI) is responsible for reviewing the draft and
    creating rows through the existing profile/skills/learning/self-pr
    endpoints."""
    data = file.file.read()
    if len(data) > MAX_FILE_SIZE_BYTES:
        raise HTTPException(status_code=400, detail=f"{file.filename}: file exceeds 10MB limit")

    ext = Path(file.filename or "").suffix.lower()
    if ext not in (".docx", ".pdf"):
        raise HTTPException(status_code=400, detail="Only .docx and .pdf files are supported")

    if ext == ".pdf":
        text = _pdf_text(data)
    else:
        file.file.seek(0)
        document = Document(file.file)
        text = _document_text(document)
    if not text.strip():
        raise HTTPException(status_code=400, detail="No text content found in the uploaded file")

    settings = session.get(Settings, 1)
    return llm.extract_resume(text, settings)


class LinkSkillIn(BaseModel):
    project_id: str
    name: str
    category: str = "未分類"


@router.post("/link-skill")
def link_skill(payload: LinkSkillIn, session: Session = Depends(get_session)) -> Skill:
    """Links a resume-extracted skill to the project it came from, via a
    SkillLink on that project's own evidence_id — the same evidence_id
    app/routers/graph.py looks for when drawing a Skill->Project edge.
    Deliberately independent of Settings.skill_extraction_enabled: that
    flag gates *passive* extraction from free text (quick updates, a
    project's own description on create), but this is a human confirming
    an already-reviewed, already-extracted skill, not a new LLM call."""
    project = session.get(Project, payload.project_id)
    if project is None or project.evidence_id is None:
        raise HTTPException(status_code=404, detail="Project not found")

    skill = services.upsert_skill(session, payload.name, payload.category)
    if skill is None:
        raise HTTPException(status_code=400, detail="Skill name cannot be empty")

    existing = session.exec(
        select(SkillLink).where(SkillLink.evidence_id == project.evidence_id, SkillLink.skill_id == skill.id)
    ).first()
    if existing is None:
        session.add(SkillLink(evidence_id=project.evidence_id, skill_id=skill.id, mention_text=payload.name))
        session.commit()
    return skill
