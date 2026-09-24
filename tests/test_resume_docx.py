from io import BytesIO

from docx import Document

from app.models import Education, Employment, LearningActivity, Project, ResumeTemplate, SelfPR, Skill
from app.resume_docx import render_resume_docx


def _paragraph_text(docx_bytes: bytes) -> str:
    doc = Document(BytesIO(docx_bytes))
    return "\n".join(p.text for p in doc.paragraphs)


def _has_table(docx_bytes: bytes) -> bool:
    return len(Document(BytesIO(docx_bytes)).tables) > 0


def test_render_resume_docx_fills_tags_from_structured_data(session, tmp_path, sample_docx_bytes):
    template_path = tmp_path / "template.docx"
    template_path.write_bytes(sample_docx_bytes())

    session.add(SelfPR(content="A short pitch", is_selected=True))
    session.add(Employment(id="emp-1", company="Acme", role="Engineer"))
    session.add(Project(employment_id="emp-1", title="Widget launch", role="Lead"))
    session.add(Project(title="Side project", role="Solo"))
    session.add(Education(school="Test University", degree="B.S.", major="CS"))
    session.add(Skill(name="Python", category="技術"))
    session.add(LearningActivity(activity_type="certification", title="AWS SAA"))
    session.commit()

    template = ResumeTemplate(name="Default", file_path=str(template_path))

    text = _paragraph_text(render_resume_docx(session, template))

    assert "A short pitch" in text
    assert "Acme" in text
    assert "Widget launch" in text
    assert "Side project" in text
    assert "Test University" in text
    assert "Python" in text
    assert "AWS SAA" in text


def test_render_resume_docx_does_not_require_llm(session, tmp_path, sample_docx_bytes, monkeypatch):
    from app import llm

    def boom(*args, **kwargs):
        raise AssertionError("resume docx generation should not call the LLM")

    monkeypatch.setattr(llm, "_complete", boom)

    template_path = tmp_path / "template.docx"
    template_path.write_bytes(sample_docx_bytes())
    template = ResumeTemplate(name="Default", file_path=str(template_path))

    render_resume_docx(session, template)  # must not raise


def test_section_format_table_renders_a_native_table(session, tmp_path, sample_docx_bytes):
    template_path = tmp_path / "template.docx"
    template_path.write_bytes(sample_docx_bytes())

    session.add(Skill(name="Python", category="技術"))
    session.commit()

    bullet_template = ResumeTemplate(name="Bullet", file_path=str(template_path), section_formats="{}")
    table_template = ResumeTemplate(name="Table", file_path=str(template_path), section_formats='{"skills": "table"}')

    assert not _has_table(render_resume_docx(session, bullet_template))
    assert _has_table(render_resume_docx(session, table_template))
