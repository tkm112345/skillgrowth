from sqlmodel import select

from app import llm
from app.models import Employment

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def test_extract_returns_llm_draft_as_is(client, sample_docx_bytes, monkeypatch):
    draft = {
        "education": [],
        "employment": [
            {"company": "Acme", "department": None, "role": None, "start_date": "2020-01-01", "end_date": None}
        ],
        "projects": [],
        "skills": [{"name": "Python", "category": "言語"}],
        "certifications": [],
        "self_pr": "テスト",
    }

    def fake_extract(text, settings):
        assert "self_pr" in text  # the fixture's placeholder paragraph text reached the LLM call
        return draft

    monkeypatch.setattr(llm, "extract_resume", fake_extract)

    resp = client.post(
        "/api/resume-import/extract",
        files={"file": ("resume.docx", sample_docx_bytes(), DOCX_MIME)},
    )
    assert resp.status_code == 200
    assert resp.json() == draft


def test_extract_reads_table_content_not_just_paragraphs(client, monkeypatch):
    """A real 職務経歴書 commonly lays its period-by-period project history
    out as a Word table, not body paragraphs — python-docx's `.paragraphs`
    silently skips table content entirely. Build a .docx with the project
    history in a table (the same shape a real resume used) and confirm
    that text reaches the LLM call."""
    from io import BytesIO

    from docx import Document

    doc = Document()
    doc.add_paragraph("職務要約")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "期間"
    table.rows[0].cells[1].text = "担当業務"
    table.rows[1].cells[0].text = "2020年4月〜2021年3月"
    table.rows[1].cells[1].text = "ロボット用ソフトウェア開発"
    buf = BytesIO()
    doc.save(buf)

    captured = {}

    def fake_extract(text, settings):
        captured["text"] = text
        return {}

    monkeypatch.setattr(llm, "extract_resume", fake_extract)

    resp = client.post(
        "/api/resume-import/extract",
        files={"file": ("resume.docx", buf.getvalue(), DOCX_MIME)},
    )
    assert resp.status_code == 200
    assert "ロボット用ソフトウェア開発" in captured["text"]
    assert "2020年4月〜2021年3月" in captured["text"]


def test_extract_rejects_non_docx_extension(client):
    resp = client.post(
        "/api/resume-import/extract",
        files={"file": ("resume.doc", b"not a real docx", "application/msword")},
    )
    assert resp.status_code == 400


def test_extract_does_not_persist_anything(client, session, sample_docx_bytes, monkeypatch):
    monkeypatch.setattr(llm, "extract_resume", lambda text, settings: {"employment": [{"company": "Acme"}]})

    resp = client.post(
        "/api/resume-import/extract",
        files={"file": ("resume.docx", sample_docx_bytes(), DOCX_MIME)},
    )
    assert resp.status_code == 200
    assert session.exec(select(Employment)).first() is None
