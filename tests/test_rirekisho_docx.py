from datetime import date
from io import BytesIO

from docx import Document

from app.models import Education, Employment, LearningActivity, PersonalInfo, RirekishoTemplate, SelfPR
from app.rirekisho_docx import render_rirekisho_docx


def _paragraph_text(docx_bytes: bytes) -> str:
    doc = Document(BytesIO(docx_bytes))
    return "\n".join(p.text for p in doc.paragraphs)


def test_render_rirekisho_docx_fills_tags_from_structured_data(session, tmp_path, sample_rirekisho_docx_bytes):
    template_path = tmp_path / "template.docx"
    template_path.write_bytes(sample_rirekisho_docx_bytes())

    session.add(
        PersonalInfo(
            id=1,
            name="Taro Yamada",
            name_kana="ヤマダ タロウ",
            birthdate=date(1990, 4, 1),
            postal_code="100-0001",
            address="Tokyo",
            phone="090-0000-0000",
            email="taro@example.com",
        )
    )
    session.add(SelfPR(content="A short pitch", is_selected=True))
    session.add(Education(school="Test University", start_date=date(2009, 4, 1), end_date=date(2013, 3, 31)))
    session.add(Employment(company="Acme", start_date=date(2013, 4, 1)))
    session.add(LearningActivity(activity_type="certification", title="AWS SAA"))
    session.commit()

    template = RirekishoTemplate(name="Default", file_path=str(template_path))

    text = _paragraph_text(render_rirekisho_docx(session, template))

    assert "Taro Yamada" in text
    assert "ヤマダ タロウ" in text
    assert "1990-04-01" in text
    assert "Tokyo" in text
    assert "taro@example.com" in text
    assert "A short pitch" in text
    assert "AWS SAA" in text


def test_render_rirekisho_docx_history_table_includes_school_and_company(
    session, tmp_path, sample_rirekisho_docx_bytes
):
    template_path = tmp_path / "template.docx"
    template_path.write_bytes(sample_rirekisho_docx_bytes())

    session.add(Education(school="Test University", start_date=date(2009, 4, 1), end_date=date(2013, 3, 31)))
    session.add(Employment(company="Acme", start_date=date(2013, 4, 1)))
    session.commit()

    template = RirekishoTemplate(name="Default", file_path=str(template_path))
    doc = Document(BytesIO(render_rirekisho_docx(session, template)))
    table_text = "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)

    assert "Test University 入学" in table_text
    assert "Test University 卒業" in table_text
    assert "Acme 入社" in table_text
    assert "現在に至る" in table_text


def test_render_rirekisho_docx_does_not_require_llm(session, tmp_path, sample_rirekisho_docx_bytes, monkeypatch):
    from app import llm

    def boom(*args, **kwargs):
        raise AssertionError("rirekisho docx generation should not call the LLM")

    monkeypatch.setattr(llm, "_complete", boom)

    template_path = tmp_path / "template.docx"
    template_path.write_bytes(sample_rirekisho_docx_bytes())
    template = RirekishoTemplate(name="Default", file_path=str(template_path))

    render_rirekisho_docx(session, template)  # must not raise


def test_render_rirekisho_docx_without_personal_info_does_not_raise(session, tmp_path, sample_rirekisho_docx_bytes):
    template_path = tmp_path / "template.docx"
    template_path.write_bytes(sample_rirekisho_docx_bytes())
    template = RirekishoTemplate(name="Default", file_path=str(template_path))

    render_rirekisho_docx(session, template)  # no PersonalInfo row at all — must not raise
